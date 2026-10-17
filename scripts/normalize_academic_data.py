"""Build reviewable source metadata and structured course/timetable CSVs.

Run from the project directory with the project environment active:
    python scripts/normalize_academic_data.py

This reads the original workbooks and processing inventory. It does not change
raw files or the sheet-by-sheet CSV exports.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = PROJECT_ROOT / "data" / "raw"
META_ROOT = PROJECT_ROOT / "data" / "metadata"
OUTPUT_ROOT = PROJECT_ROOT / "data" / "processed" / "normalized"
COURSE_BOOK = RAW_ROOT / "timetables" / "Finalized courses list.xlsx"
SLOT_BOOKS = [
    RAW_ROOT / "timetables" / "AY 2026-27_Semester I_Slot System (2nd Year onwards).xlsx",
    RAW_ROOT / "timetables" / "AY 2026-27 Semester-I Slot System (B.Tech. 1st Year).xlsx",
]

CODE_RE = re.compile(r"^[A-Z]{2,4}[A-Z0-9]{4,6}$")
LTP_RE = re.compile(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)")
SLOT_RE = re.compile(r"^[A-Z][A-Z0-9]?$", re.IGNORECASE)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def canonical_code(value: Any) -> str:
    code = re.sub(r"[^A-Za-z0-9]", "", clean(value)).upper()
    return code if CODE_RE.fullmatch(code) else ""


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_updated{path.suffix}")
        with alt.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        print(f"Locked output {path.name}; saved {alt.name} instead.")
        return alt


def find_header_map(values: tuple[Any, ...] | list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(values):
        key = header_key(value)
        if not key:
            continue
        if key in {"coursecode", "code", "subjectcode"} or key.startswith("coursecode"):
            result.setdefault("course_code", index)
        if ("course" in key or "subject" in key) and ("name" in key or "title" in key):
            result.setdefault("course_name", index)
        if key == "subject":
            result.setdefault("course_name", index)
        if key == "nameofthecourse":
            result.setdefault("course_name", index)
        if "ltp" in key or "ltpstructure" in key:
            result.setdefault("ltp", index)
        if "credit" in key:
            result.setdefault("credits", index)
        if "slot" in key and "lab" not in key:
            result.setdefault("slot", index)
        if "instructor" in key or "faculty" in key and "name" in key or "coordinator" in key:
            result.setdefault("instructor", index)
        if "coinstructor" in key:
            result.setdefault("co_instructors", index)
        if "category" in key or key in {"type", "coursetype"}:
            result.setdefault("category", index)
        if "remark" in key or "comment" in key:
            result.setdefault("remarks", index)
        if "programme" in key or "program" in key or "offeredfor" in key or "batch" in key or "year" == key:
            result.setdefault("program_or_batch", index)
        if "department" in key:
            result.setdefault("department", index)
        if "registered" in key or "enrolled" in key:
            result.setdefault("enrollment", index)
        if "classroom" in key or "room" == key:
            result.setdefault("classroom", index)
    return result


def value_at(row: tuple[Any, ...] | list[Any], columns: dict[str, int], key: str) -> str:
    index = columns.get(key)
    return clean(row[index]) if index is not None and index < len(row) else ""


def parse_ltp_and_credit(ltp_value: str, credit_value: str) -> tuple[str, str]:
    match = LTP_RE.search(ltp_value)
    ltp = "-".join(match.groups()) if match else ltp_value.strip("[]")
    credit = credit_value
    if match and not credit:
        remainder = ltp_value[match.end():]
        credit_match = re.search(r"\(\s*(\d+(?:\.\d+)?)\s*\)", remainder)
        if credit_match:
            credit = credit_match.group(1)
    return ltp, credit


def collapse_exact_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Collapse identical normalized rows while retaining all workbook row refs."""
    result: list[dict[str, Any]] = []
    by_content: dict[tuple[tuple[str, str], ...], dict[str, Any]] = {}
    for row in rows:
        content = tuple(sorted((key, clean(value)) for key, value in row.items() if key != "source_row"))
        existing = by_content.get(content)
        if existing is None:
            row = dict(row)
            row["source_rows"] = clean(row.get("source_row", ""))
            by_content[content] = row
            result.append(row)
        else:
            refs = [x for x in existing["source_rows"].split(";") if x]
            current = clean(row.get("source_row", ""))
            if current and current not in refs:
                refs.append(current)
            existing["source_rows"] = ";".join(refs)
    return result


def build_course_offerings() -> list[dict[str, Any]]:
    if not COURSE_BOOK.is_file():
        print(f"Course workbook missing: {COURSE_BOOK}")
        return []
    workbook = load_workbook(COURSE_BOOK, read_only=True, data_only=True)
    offerings: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        active_map: dict[str, int] | None = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            mapping = find_header_map(row)
            if "course_code" in mapping and "course_name" in mapping:
                active_map = mapping
                continue
            if not active_map:
                continue
            code = canonical_code(value_at(row, active_map, "course_code"))
            name = value_at(row, active_map, "course_name")
            if not code or not name:
                continue
            ltp, credits = parse_ltp_and_credit(value_at(row, active_map, "ltp"), value_at(row, active_map, "credits"))
            suffix = re.match(r"^[A-Z]{2,4}(\d)", code)
            offerings.append({
                "course_code": code,
                "course_name": name,
                "department_sheet": sheet.title,
                "program_or_batch": value_at(row, active_map, "program_or_batch"),
                "category": value_at(row, active_map, "category"),
                "credits": credits,
                "ltp": ltp,
                "slot": value_at(row, active_map, "slot"),
                "instructor": value_at(row, active_map, "instructor"),
                "co_instructors": value_at(row, active_map, "co_instructors"),
                "remarks": value_at(row, active_map, "remarks"),
                "course_level_digit": suffix.group(1) if suffix else "",
                "source_file": COURSE_BOOK.relative_to(PROJECT_ROOT).as_posix(),
                "source_sheet": sheet.title,
                "source_row": row_number,
            })
    workbook.close()
    return collapse_exact_rows(offerings)


def build_classroom_allocation() -> list[dict[str, Any]]:
    rows_out: list[dict[str, Any]] = []
    for path in SLOT_BOOKS:
        if not path.is_file():
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            if "classroom" not in sheet.title.casefold() and "allocation" not in sheet.title.casefold():
                continue
            active: dict[str, int] | None = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                mapping = find_header_map(row)
                # In the allocation sheet the first, unnamed column is Course Code.
                if "course_code" not in mapping and any("subject" in header_key(v) for v in row):
                    mapping["course_code"] = 0
                if "course_code" in mapping and ("course_name" in mapping or "slot" in mapping):
                    active = mapping
                    continue
                if not active:
                    continue
                code_index = active.get("course_code", 0)
                code = canonical_code(row[code_index] if code_index < len(row) else "")
                if not code:
                    continue
                ltp, credits = parse_ltp_and_credit(value_at(row, active, "ltp"), value_at(row, active, "credits"))
                rows_out.append({
                    "course_code": code,
                    "course_name": value_at(row, active, "course_name"),
                    "ltp": ltp,
                    "credits": credits,
                    "slot": value_at(row, active, "slot"),
                    "department": value_at(row, active, "department"),
                    "instructor": value_at(row, active, "instructor"),
                    "co_instructors": value_at(row, active, "co_instructors"),
                    "enrollment": value_at(row, active, "enrollment"),
                    "classroom": value_at(row, active, "classroom"),
                    "source_file": path.relative_to(PROJECT_ROOT).as_posix(),
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                })
        workbook.close()
    return collapse_exact_rows(rows_out)


def build_slot_times() -> list[dict[str, Any]]:
    rows_out: list[dict[str, Any]] = []
    path = SLOT_BOOKS[0]
    if not path.is_file():
        return rows_out
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        if "slot system" not in sheet.title.casefold():
            continue
        values = list(sheet.iter_rows(values_only=True))
        day_row_index = None
        day_columns: list[tuple[str, int]] = []
        for row_index, row in enumerate(values):
            found = []
            for col, cell in enumerate(row):
                day = clean(cell).title()
                if day in {"Mon", "Tue", "Wed", "Thu", "Fri"}:
                    found.append((day, col))
            if len(found) >= 4:
                day_row_index, day_columns = row_index, found
                break
        if day_row_index is None:
            continue
        for row_index in range(day_row_index + 1, len(values)):
            row = values[row_index]
            time_value = clean(row[0]) if row else ""
            if not re.match(r"^\d{1,2}\s*[:–-]", time_value):
                continue
            for day_index, (day, col) in enumerate(day_columns):
                stop = day_columns[day_index + 1][1] if day_index + 1 < len(day_columns) else min(len(row), col + 2)
                for cell in row[col:stop]:
                    code = clean(cell).upper()
                    if SLOT_RE.fullmatch(code):
                        rows_out.append({
                            "slot": code,
                            "day": day,
                            "time": time_value,
                            "source_file": path.relative_to(PROJECT_ROOT).as_posix(),
                            "source_sheet": sheet.title,
                            "source_row": row_index + 1,
                        })
    workbook.close()
    # Preserve unique day/time occurrences while keeping separate repeated slots.
    unique: dict[tuple[str, str, str], dict[str, Any]] = {}
    for row in rows_out:
        unique[(row["slot"], row["day"], row["time"])] = row
    return list(unique.values())


def topic_tags(relative_path: str) -> str:
    text = relative_path.casefold()
    terms = [
        ("attendance", "attendance"), ("exam", "examinations"), ("registration", "registration"),
        ("curriculum", "curriculum"), ("course", "courses"), ("slot", "timetable"),
        ("timetable", "timetable"), ("branch change", "branch-change"), ("minor", "minors"),
        ("specialization", "specializations"), ("calendar", "academic-calendar"),
        ("faculty", "faculty-contacts"), ("people", "faculty-contacts"),
        ("fee", "fees"), ("credit", "credits"), ("regulation", "regulations"),
    ]
    return ";".join(tag for key, tag in terms if key in text)


def build_source_register() -> list[dict[str, Any]]:
    inventory_path = META_ROOT / "raw_inventory.csv"
    if not inventory_path.is_file():
        print(f"Run process_raw_data.py first; inventory not found: {inventory_path}")
        return []
    with inventory_path.open(encoding="utf-8-sig", newline="") as stream:
        inventory = list(csv.DictReader(stream))
    canonical = [row for row in inventory if Path(row["relative_path"]).suffix.lower() != ".txt"]
    canonical.sort(key=lambda row: row["relative_path"].casefold())
    ids = {row["relative_path"]: f"DOC{index:04d}" for index, row in enumerate(canonical, start=1)}
    out = []
    for row in inventory:
        relative = row["relative_path"]
        file_path = Path(relative)
        parent = file_path.parts[0] if file_path.parts else ""
        sidecar_of = ""
        citation_id = ids.get(relative, "")
        if file_path.suffix.lower() == ".txt":
            pdf_path = file_path.with_suffix(".pdf").as_posix()
            if pdf_path in ids:
                sidecar_of = pdf_path
                citation_id = ids[pdf_path]
        out.append({
            "citation_id": citation_id,
            "title": file_path.stem,
            "raw_path": relative,
            "processed_path": row.get("output_path", ""),
            "source_category": parent,
            "document_type": file_path.suffix.lower().lstrip("."),
            "processing_status": row.get("processing_status", ""),
            "topic_tags": topic_tags(relative),
            "source_authority": "IIT Jodhpur (verify)",
            "source_url": "",
            "publication_date": "",
            "effective_from": "",
            "effective_to": "",
            "applicable_program_batch": "",
            "supersedes": "",
            "sidecar_of": sidecar_of,
            "include_in_rag": "yes" if row.get("processing_status") in {"processed", "processed_ocr_text"} else "no",
            "review_status": "needs_source_and_scope_review",
        })
    return out


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    offerings = build_course_offerings()
    allocations = build_classroom_allocation()
    slots = build_slot_times()
    register = build_source_register()
    paths = [
        write_csv(OUTPUT_ROOT / "course_offerings.csv", offerings, [
            "course_code", "course_name", "department_sheet", "program_or_batch", "category", "credits", "ltp",
            "slot", "instructor", "co_instructors", "remarks", "course_level_digit", "source_file", "source_sheet", "source_row", "source_rows",
        ]),
        write_csv(OUTPUT_ROOT / "classroom_allocation.csv", allocations, [
            "course_code", "course_name", "ltp", "credits", "slot", "department", "instructor", "co_instructors",
            "enrollment", "classroom", "source_file", "source_sheet", "source_row", "source_rows",
        ]),
        write_csv(OUTPUT_ROOT / "slot_times.csv", slots, ["slot", "day", "time", "source_file", "source_sheet", "source_row"]),
        write_csv(META_ROOT / "source_register.csv", register, [
            "citation_id", "title", "raw_path", "processed_path", "source_category", "document_type", "processing_status",
            "topic_tags", "source_authority", "source_url", "publication_date", "effective_from", "effective_to",
            "applicable_program_batch", "supersedes", "sidecar_of", "include_in_rag", "review_status",
        ]),
    ]
    print(f"Course offerings normalized: {len(offerings)}")
    print(f"Classroom allocation rows: {len(allocations)}")
    print(f"Slot/day/time rows: {len(slots)}")
    print(f"Source register rows: {len(register)}")
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
