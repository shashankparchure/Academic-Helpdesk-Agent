"""Extract raw Academic Helpdesk files into searchable text and per-sheet CSVs.

Run from the project directory with its virtual environment active:
    python scripts/process_raw_data.py

Raw files are never modified. Legacy .doc files are inventoried and reported as
unsupported; convert those to .docx or PDF before processing if their contents
are needed. Scanned PDFs with no extractable text are flagged for OCR.
"""

from __future__ import annotations

import csv
import hashlib
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = PROJECT_ROOT / "data" / "raw"
TEXT_ROOT = PROJECT_ROOT / "data" / "processed" / "extracted_text"
SHEET_ROOT = PROJECT_ROOT / "data" / "processed" / "spreadsheets"
META_ROOT = PROJECT_ROOT / "data" / "metadata"


def slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return value.strip("._") or "unnamed"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def exclusion_reason(path: Path) -> str:
    name = path.name.casefold()
    if "salary" in name:
        return "Excluded: personal salary data; not needed for academic helpdesk."
    if "private visits abroad" in name or "higher education at iit jodhpur" in name:
        return "Excluded: HR permission form, outside the academic helpdesk scope."
    return ""


def extract_pdf(path: Path) -> tuple[str, int, int]:
    reader = PdfReader(str(path))
    pages = []
    nonempty = 0
    for number, page in enumerate(reader.pages, start=1):
        body = (page.extract_text() or "").strip()
        if body:
            nonempty += 1
        pages.append(f"## Page {number}\n\n{body}")
    return "\n\n".join(pages), len(reader.pages), nonempty


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table_no, table in enumerate(doc.tables, start=1):
        parts.append(f"\n## Table {table_no}")
        for row in table.rows:
            parts.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
    return "\n\n".join(parts)


def read_text_file(path: Path) -> tuple[str, bool]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    markers = ("à¤", "à¥", "â€“", "â€™", "Ã")
    repaired = False
    if any(marker in text for marker in markers):
        try:
            candidate = text.encode("latin-1").decode("utf-8")
            if sum(candidate.count(marker) for marker in markers) < sum(text.count(marker) for marker in markers):
                text = candidate
                repaired = True
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return text.strip(), repaired


def export_workbook(path: Path) -> tuple[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    output_dir = SHEET_ROOT / slug(path.parent.name) / slug(path.stem)
    output_dir.mkdir(parents=True, exist_ok=True)
    markdown = []
    for sheet in workbook.worksheets:
        sheet_path = output_dir / f"{slug(sheet.title)}.csv"
        row_count = 0
        with sheet_path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.writer(stream)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
                row_count += 1
        markdown.append(f"## Sheet: {sheet.title}\n\nCSV: {sheet_path.relative_to(PROJECT_ROOT).as_posix()}\n\nRows exported: {row_count}")
    workbook.close()
    return "\n\n".join(markdown), len(markdown)


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return path
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        alternate = path.with_name(f"{path.stem}_updated_{stamp}{path.suffix}")
        with alternate.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        print(f"Could not overwrite locked file: {path.name}; saved updated copy as {alternate.name}")
        return alternate


def main() -> None:
    if not RAW_ROOT.exists():
        raise SystemExit(f"Raw data folder not found: {RAW_ROOT}")
    TEXT_ROOT.mkdir(parents=True, exist_ok=True)
    SHEET_ROOT.mkdir(parents=True, exist_ok=True)
    META_ROOT.mkdir(parents=True, exist_ok=True)

    inventory = []
    report = []
    hashes: dict[str, list[str]] = {}
    files = sorted(p for p in RAW_ROOT.rglob("*") if p.is_file())
    text_sidecars = {p.with_suffix(".txt") for p in files if p.suffix.lower() == ".pdf" and p.with_suffix(".txt").is_file()}

    for path in files:
        relative = path.relative_to(RAW_ROOT).as_posix()
        digest = sha256(path)
        hashes.setdefault(digest, []).append(relative)
        row = {
            "relative_path": relative,
            "extension": path.suffix.lower(),
            "bytes": path.stat().st_size,
            "sha256": digest,
            "processing_status": "pending",
            "output_path": "",
            "pages_or_sheets": "",
            "extractable_pages": "",
            "notes": "",
        }
        inventory.append(row.copy())

        canonical = hashes[digest][0]
        if relative != canonical:
            row.update(
                processing_status="duplicate",
                notes=f"Byte-identical to {canonical}; retained in inventory, skipped duplicate extraction.",
            )
            report.append(row)
            continue

        reason = exclusion_reason(path)
        if reason:
            row.update(processing_status="excluded", notes=reason)
        else:
            try:
                if path.suffix.lower() == ".pdf":
                    sidecar = path.with_suffix(".txt")
                    if sidecar.is_file():
                        text, repaired = read_text_file(sidecar)
                        out = TEXT_ROOT / slug(path.parent.name) / f"{slug(path.stem)}.md"
                        out.parent.mkdir(parents=True, exist_ok=True)
                        note = f"Used sibling OCR text file: {sidecar.name}."
                        if repaired:
                            note += " Repaired UTF-8 text that had been misread as Latin-1."
                        out.write_text(f"# Source: {path.name}\n\nRaw path: {relative}\n\nOCR text source: {sidecar.name}\n\n{text}\n", encoding="utf-8")
                        row.update(processing_status="processed_ocr_text", output_path=out.relative_to(PROJECT_ROOT).as_posix(), notes=note)
                    else:
                        text, page_count, extracted_pages = extract_pdf(path)
                        if extracted_pages == 0:
                            row.update(
                                processing_status="needs_ocr",
                                pages_or_sheets=page_count,
                                extractable_pages=0,
                                notes="No extractable text detected; likely scanned PDF.",
                            )
                        else:
                            row.update(pages_or_sheets=page_count, extractable_pages=extracted_pages)
                            if extracted_pages < page_count:
                                row["notes"] = "Some pages have no extractable text; inspect for scans or image-only pages."
                            out = TEXT_ROOT / slug(path.parent.name) / f"{slug(path.stem)}.md"
                            out.parent.mkdir(parents=True, exist_ok=True)
                            out.write_text(f"# Source: {path.name}\n\nRaw path: {relative}\n\n{text}\n", encoding="utf-8")
                            row.update(processing_status="processed", output_path=out.relative_to(PROJECT_ROOT).as_posix())
                elif path.suffix.lower() == ".docx":
                    text = extract_docx(path)
                    out = TEXT_ROOT / slug(path.parent.name) / f"{slug(path.stem)}.md"
                    out.parent.mkdir(parents=True, exist_ok=True)
                    out.write_text(f"# Source: {path.name}\n\nRaw path: {relative}\n\n{text}\n", encoding="utf-8")
                    row.update(processing_status="processed", output_path=out.relative_to(PROJECT_ROOT).as_posix())
                elif path.suffix.lower() in {".xlsx", ".xlsm"}:
                    text, sheet_count = export_workbook(path)
                    out = TEXT_ROOT / slug(path.parent.name) / f"{slug(path.stem)}.md"
                    out.parent.mkdir(parents=True, exist_ok=True)
                    out.write_text(f"# Workbook: {path.name}\n\nRaw path: {relative}\n\n{text}\n", encoding="utf-8")
                    row.update(processing_status="processed", pages_or_sheets=sheet_count, output_path=out.relative_to(PROJECT_ROOT).as_posix())
                elif path.suffix.lower() == ".txt":
                    if path in text_sidecars:
                        source_pdf = path.with_suffix(".pdf")
                        out = TEXT_ROOT / slug(path.parent.name) / f"{slug(source_pdf.stem)}.md"
                        row.update(
                            processing_status="sidecar",
                            output_path=out.relative_to(PROJECT_ROOT).as_posix(),
                            notes=f"OCR text is incorporated into the matching PDF source: {source_pdf.name}.",
                        )
                    else:
                        text, repaired = read_text_file(path)
                        out = TEXT_ROOT / slug(path.parent.name) / f"{slug(path.stem)}.md"
                        out.parent.mkdir(parents=True, exist_ok=True)
                        out.write_text(f"# Source: {path.name}\n\nRaw path: {relative}\n\n{text}\n", encoding="utf-8")
                        notes = "Repaired UTF-8 text that had been misread as Latin-1." if repaired else ""
                        row.update(processing_status="processed", output_path=out.relative_to(PROJECT_ROOT).as_posix(), notes=notes)
                elif path.suffix.lower() == ".doc":
                    row.update(processing_status="unsupported", notes="Legacy .doc format; convert to .docx or PDF, then rerun.")
                else:
                    row.update(processing_status="unsupported", notes="No extractor configured for this file type.")
            except Exception as exc:
                row.update(processing_status="error", notes=f"{type(exc).__name__}: {exc}")
        report.append(row)

    report_by_path = {row["relative_path"]: row for row in report}
    for source in inventory:
        processed = report_by_path[source["relative_path"]]
        for key in ("processing_status", "output_path", "pages_or_sheets", "extractable_pages", "notes"):
            source[key] = processed[key]

    columns = list(inventory[0].keys()) if inventory else []
    inventory_path = write_csv(META_ROOT / "raw_inventory.csv", inventory, columns)
    report_path = write_csv(META_ROOT / "processing_report.csv", report, list(report[0].keys()) if report else columns)
    duplicate_rows = [
        {"sha256": digest, "relative_paths": " | ".join(paths), "count": len(paths)}
        for digest, paths in hashes.items() if len(paths) > 1
    ]
    duplicates_path = write_csv(META_ROOT / "duplicate_files.csv", duplicate_rows, ["sha256", "relative_paths", "count"])

    counts: dict[str, int] = {}
    for row in report:
        counts[row["processing_status"]] = counts.get(row["processing_status"], 0) + 1
    print(f"Files inventoried: {len(files)}")
    print("Processing statuses: " + ", ".join(f"{key}={value}" for key, value in sorted(counts.items())))
    print(f"Inventory: {inventory_path}")
    print(f"Report:    {report_path}")
    print(f"Duplicates:{duplicates_path}")
    print(f"Sheets:    {SHEET_ROOT}")
    print(f"Text:      {TEXT_ROOT}")


if __name__ == "__main__":
    main()
